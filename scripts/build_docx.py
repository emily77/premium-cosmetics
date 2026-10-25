# -*- coding: utf-8 -*-
"""
build_docx.py — 產生「頂級化妝品產品線」彩色 Word 分析報告

資料來源（單一真相）：
  data/finance_model.json   財務模型（投資、化工廠、AI 團隊、定價、損益、回收、瀑布）
  data/cost_report.csv      14 支 SKU 成本彙總
  data/cost_detail.csv      248 筆原料明細
  data/products.json        產品主檔
  data/ingredient_prices.json  抓價結果

輸出：
  docs/charts/*.png
  docs/07_投資與獲利分析報告.docx

用法：
  python scripts\\build_docx.py
"""
from __future__ import annotations

import csv
import json
import sys
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent.parent
DATA = BASE / "data"
DOCS = BASE / "docs"
CHARTS = DOCS / "charts"
OUT_DOCX = DOCS / "07_投資與獲利分析報告.docx"

# ── 品牌色 ────────────────────────────────────────────────────────────────
NAVY = "1F3864"
NAVY2 = "2E5496"
GOLD = "C9A227"
GOLD_L = "F7EDCB"
GREY_L = "F2F2F2"
GREY_M = "D9D9D9"
GREY_D = "595959"
RED = "C00000"
RED_L = "FBE4E4"
ORG = "ED7D31"
ORG_L = "FDEADA"
GRN = "548235"
GRN_L = "E2EFDA"
BLU_L = "DEEAF6"
WHITE = "FFFFFF"

BODY_FONT = "Microsoft JhengHei"
HEAD_FONT = "Microsoft JhengHei"

PRIORITY_STYLE = {
    "P0": (RED_L, RED),
    "P1": (ORG_L, ORG),
    "P2": (GREY_L, GREY_D),
}


# ── 小工具 ────────────────────────────────────────────────────────────────
def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def n(v, dec: int = 0) -> str:
    try:
        return f"{float(v):,.{dec}f}"
    except (TypeError, ValueError):
        return str(v)


def pct(v, dec: int = 1) -> str:
    try:
        return f"{float(v):.{dec}f}%"
    except (TypeError, ValueError):
        return str(v)


# OOXML 子元素必須照 schema 順序排列，否則 Word 可能報「檔案已損毀」
PPR_ORDER = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
    "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs",
    "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
    "w:sectPr", "w:pPrChange",
]
TBLPR_ORDER = [
    "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
    "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
    "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
    "w:tblCaption", "w:tblDescription",
]
TCPR_ORDER = [
    "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd",
    "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
    "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange",
]
TRPR_ORDER = [
    "w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore", "w:wAfter",
    "w:cantSplit", "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden",
    "w:ins", "w:del", "w:trPrChange",
]


def _order_index(order, tag):
    for i, t in enumerate(order):
        if qn(t) == tag:
            return i
    return len(order)


def insert_ordered(parent, el, order):
    """依 schema 順序插入子元素。"""
    my = _order_index(order, el.tag)
    for child in list(parent):
        if _order_index(order, child.tag) > my:
            child.addprevious(el)
            return
    parent.append(el)


def shade(cell, hexcolor: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    insert_ordered(tcPr, shd, TCPR_ORDER)


def table_borders(table, color: str = "BFBFBF", sz: int = 4) -> None:
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    insert_ordered(tblPr, borders, TBLPR_ORDER)


def no_split_rows(table) -> None:
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            insert_ordered(trPr, OxmlElement("w:cantSplit"), TRPR_ORDER)


def add_pbdr(p, color: str = GOLD, sz: int = 12, space: int = 4) -> None:
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:pBdr")):
        pPr.remove(old)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    insert_ordered(pPr, pbdr, PPR_ORDER)


def set_run(run, size=10, bold=False, color=None, font=BODY_FONT, italic=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:eastAsia"), font)
    if color:
        run.font.color.rgb = rgb(color)
    return run


def set_cell(cell, text, bold=False, color=None, size=9, align="left", bg=None,
             font=BODY_FONT, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    pf = p.paragraph_format
    pf.space_before = Pt(1.5)
    pf.space_after = Pt(1.5)
    pf.line_spacing = 1.0
    set_run(p.add_run(str(text)), size=size, bold=bold, color=color, font=font, italic=italic)
    if bg:
        shade(cell, bg)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, widths, aligns=None, header_bg=NAVY,
              size=9, row_bg=None, row_fg=None, bold_rows=(), header_size=9,
              zebra=False, first_col_bold=False, keep_together=None):
    """建立彩色表格。widths 單位 cm。"""
    ncol = len(headers)
    aligns = aligns or ["left"] * ncol
    table = doc.add_table(rows=1, cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_borders(table)

    hdr = table.rows[0]
    hdr.height = Cm(0.72)
    for i, h in enumerate(headers):
        set_cell(hdr.cells[i], h, bold=True, color=WHITE, size=header_size,
                 align="center", bg=header_bg, font=HEAD_FONT)

    for ri, row in enumerate(rows):
        tr = table.add_row()
        bg = row_bg[ri] if row_bg else (GREY_L if (zebra and ri % 2 == 1) else None)
        fg = row_fg[ri] if row_fg else None
        for ci, val in enumerate(row):
            bold = (ri in bold_rows) or (first_col_bold and ci == 0)
            set_cell(tr.cells[ci], val, bold=bold, color=fg, size=size,
                     align=aligns[ci], bg=bg)

    for row in table.rows:
        for ci, w in enumerate(widths):
            row.cells[ci].width = Cm(w)

    no_split_rows(table)
    if keep_together is None:
        keep_together = len(rows) <= 15
    if keep_together:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
    return table


def h1(doc, text, num=None, new_page=False):
    p = doc.add_paragraph(style="Heading 1")
    pf = p.paragraph_format
    pf.space_before = Pt(20)
    pf.space_after = Pt(10)
    pf.keep_with_next = True
    if new_page:
        pf.page_break_before = True
    if num:
        set_run(p.add_run(f"{num}  "), size=19, bold=True, color=GOLD, font=HEAD_FONT)
    set_run(p.add_run(text), size=19, bold=True, color=NAVY, font=HEAD_FONT)
    add_pbdr(p, GOLD, 12, 4)
    return p


def h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    pf = p.paragraph_format
    pf.space_before = Pt(14)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    set_run(p.add_run("▍"), size=12, bold=True, color=GOLD, font=HEAD_FONT)
    set_run(p.add_run(text), size=13, bold=True, color=NAVY2, font=HEAD_FONT)
    return p


def h1_plain(doc, text, new_page=False):
    """不進目錄的大標題。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(20)
    pf.space_after = Pt(10)
    pf.keep_with_next = True
    if new_page:
        pf.page_break_before = True
    set_run(p.add_run(text), size=19, bold=True, color=NAVY, font=HEAD_FONT)
    add_pbdr(p, GOLD, 12, 4)
    return p


def h2_plain(doc, text):
    """不進目錄的小標題。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(14)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    set_run(p.add_run("▍"), size=12, bold=True, color=GOLD, font=HEAD_FONT)
    set_run(p.add_run(text), size=13, bold=True, color=NAVY2, font=HEAD_FONT)
    return p


def h3(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(4)
    pf.keep_with_next = True
    set_run(p.add_run(text), size=11, bold=True, color=NAVY, font=HEAD_FONT)
    return p


def para(doc, text, size=10, color=None, bold=False, align="left", space_after=6,
         indent=0.0, italic=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.45
    if indent:
        pf.left_indent = Cm(indent)
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    set_run(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p


def bullet(doc, text, size=10, color=None, level=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.6 + 0.6 * level)
    pf.first_line_indent = Cm(-0.45)
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.4
    set_run(p.add_run("• " if level == 0 else "– "), size=size, bold=True, color=GOLD)
    set_run(p.add_run(text), size=size, color=color)
    return p


def callout(doc, title, text, bg=BLU_L, bar=NAVY2, title_color=NAVY):
    """單格底色提示框。"""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t, color=bar, sz=6)
    cell = t.rows[0].cells[0]
    cell.width = Cm(16.4)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(title), size=10, bold=True, color=title_color, font=HEAD_FONT)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(4)
    p2.paragraph_format.line_spacing = 1.4
    set_run(p2.add_run(text), size=9.5, color="333333")
    shade(cell, bg)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    set_run(p.add_run(text), size=8.5, color=GREY_D, italic=True)
    return p


def add_picture(doc, path: Path, width_cm=16.0):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_field(paragraph, instr: str, size=9, color=GREY_D):
    run = paragraph.add_run()
    set_run(run, size=size, color=color)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_txt = OxmlElement("w:t")
    fld_txt.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_el)
    run._r.append(fld_sep)
    run._r.append(fld_txt)
    run._r.append(fld_end)


# ── 圖表 ──────────────────────────────────────────────────────────────────
def setup_mpl():
    plt.rcParams["font.sans-serif"] = ["Microsoft JhengHei", "Microsoft YaHei", "SimHei"]
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 200
    plt.rcParams["savefig.bbox"] = "tight"
    plt.rcParams["axes.edgecolor"] = "#BFBFBF"
    plt.rcParams["axes.labelcolor"] = "#404040"
    plt.rcParams["xtick.color"] = "#404040"
    plt.rcParams["ytick.color"] = "#404040"


def chart_cost_structure(cost: dict) -> Path:
    codes = list(cost.keys())
    labels = [f"{c}\n{cost[c]['name_zh'][:6]}" for c in codes]
    raw = [cost[c]["raw"] for c in codes]
    pack = [cost[c]["pack"] for c in codes]
    fill = [cost[c]["fill"] for c in codes]
    qc = [cost[c]["qc"] for c in codes]

    fig, ax = plt.subplots(figsize=(9.2, 4.6))
    x = range(len(codes))
    b1 = ax.bar(x, raw, color="#2E5496", label="料體")
    b2 = ax.bar(x, pack, bottom=raw, color="#C9A227", label="包材")
    b3 = ax.bar(x, fill, bottom=[a + b for a, b in zip(raw, pack)], color="#ED7D31", label="充填")
    b4 = ax.bar(x, qc, bottom=[a + b + c for a, b, c in zip(raw, pack, fill)],
                color="#A6A6A6", label="檢驗")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels, fontsize=7)
    ax.set_ylabel("TWD／支（出廠前）", fontsize=9)
    ax.set_title("14 支 SKU 成本結構（料體／包材／充填／檢驗）", fontsize=11,
                 color="#1F3864", fontweight="bold", pad=12)
    ax.legend(fontsize=8, ncol=4, frameon=False, loc="upper center",
              bbox_to_anchor=(0.5, -0.16))
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color="#E7E7E7", linewidth=0.8)
    ax.set_axisbelow(True)
    fig.savefig(CHARTS / "01_cost_structure.png")
    plt.close(fig)
    return CHARTS / "01_cost_structure.png"


def chart_investment(by_cat: list) -> Path:
    labels = [r["category"] for r in by_cat]
    vals = [r["amount_twd"] for r in by_cat]
    colors = ["#1F3864", "#2E5496", "#C9A227", "#ED7D31", "#548235", "#7F7F7F", "#A6A6A6"]

    fig, ax = plt.subplots(figsize=(8.4, 4.4))
    wedges, _t, autot = ax.pie(
        vals, colors=colors[: len(vals)], startangle=90, counterclock=False,
        autopct=lambda p: f"{p:.1f}%", pctdistance=0.78,
        wedgeprops=dict(width=0.42, edgecolor="white", linewidth=1.6),
        textprops=dict(fontsize=8, color="white", fontweight="bold"),
    )
    total = sum(vals)
    ax.text(0, 0.10, "一次性投資", ha="center", va="center", fontsize=9, color="#595959")
    ax.text(0, -0.10, f"NT$ {total/10000:,.0f} 萬", ha="center", va="center",
            fontsize=13, color="#1F3864", fontweight="bold")
    ax.legend(
        [f"{l}　{n(v/10000,1)} 萬" for l, v in zip(labels, vals)],
        fontsize=8, frameon=False, loc="center left", bbox_to_anchor=(1.0, 0.5),
    )
    ax.set_title("一次性投資結構（7 大類）", fontsize=11, color="#1F3864",
                 fontweight="bold", pad=10)
    fig.savefig(CHARTS / "02_investment.png")
    plt.close(fig)
    return CHARTS / "02_investment.png"


def chart_ai(ai: dict) -> Path:
    rows = ai["rows"]
    labels = [f"{r['code']} {r['name']}" for r in rows]
    human = [r["human_cost_twd"] for r in rows]
    aic = [r["ai_cost_twd"] for r in rows]

    fig, ax = plt.subplots(figsize=(8.8, 3.9))
    y = range(len(rows))
    ax.barh([i + 0.2 for i in y], human, height=0.36, color="#1F3864", label="純人力月成本")
    ax.barh([i - 0.2 for i in y], aic, height=0.36, color="#C9A227", label="AI Agent 月成本")
    ax.set_yticks(list(y))
    ax.set_yticklabels(labels, fontsize=8.5)
    ax.invert_yaxis()
    ax.set_xlabel("TWD／月", fontsize=9)
    ax.set_title("六個值班入口：人力成本 vs AI Agent 成本", fontsize=11,
                 color="#1F3864", fontweight="bold", pad=10)
    ax.legend(fontsize=8, frameon=False, loc="lower right")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", color="#E7E7E7", linewidth=0.8)
    ax.set_axisbelow(True)
    for i, v in enumerate(human):
        ax.text(v + 400, i + 0.2, f"{v:,.0f}", va="center", fontsize=7, color="#1F3864")
    fig.savefig(CHARTS / "03_ai_team.png")
    plt.close(fig)
    return CHARTS / "03_ai_team.png"


def chart_waterfall(wf: dict) -> Path:
    items = wf["items"]
    labels = [i["name"] for i in items]
    vals = [i["amount"] for i in items]
    colors = []
    for it in items:
        if it["name"].startswith("品牌稅後淨利"):
            colors.append("#548235")
        elif it["name"] in ("營所稅",):
            colors.append("#7F7F7F")
        else:
            colors.append("#2E5496")

    fig, ax = plt.subplots(figsize=(8.8, 4.2))
    y = range(len(items))
    ax.barh(list(y), vals, color=colors, height=0.62)
    ax.set_yticks(list(y))
    ax.set_yticklabels(labels, fontsize=8.5)
    ax.invert_yaxis()
    ax.set_xlabel("TWD／支（以建議零售價計算）", fontsize=9)
    ax.set_title(f"價值鏈拆解：{wf['sku_code']} {wf['name_zh']}（零售價 NT$ {n(wf['retail_price'])}）",
                 fontsize=10.5, color="#1F3864", fontweight="bold", pad=10)
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", color="#E7E7E7", linewidth=0.8)
    ax.set_axisbelow(True)
    for i, it in enumerate(items):
        ax.text(it["amount"] + 2, i, f"{it['amount']:,.1f}（{it['pct']:.1f}%）",
                va="center", fontsize=7.2, color="#404040")
    ax.set_xlim(0, max(vals) * 1.32)
    fig.savefig(CHARTS / "04_waterfall.png")
    plt.close(fig)
    return CHARTS / "04_waterfall.png"


def chart_pricing(pricing: dict) -> Path:
    orig = pricing["original"]
    sug = pricing["suggested"]
    labels = [f"{o['sku_code']}\n{o['name_zh'][:6]}" for o in orig]
    om = [o["multiple"] for o in orig]
    sm = [s["multiple"] for s in sug]

    fig, ax = plt.subplots(figsize=(9.2, 3.9))
    x = range(len(labels))
    ax.bar([i - 0.2 for i in x], om, width=0.38, color="#C00000", label="原提案價倍率")
    ax.bar([i + 0.2 for i in x], sm, width=0.38, color="#548235",
           label=f"建議價倍率（{pricing['target_multiple']:.0f}×）")
    ax.axhline(12, color="#C9A227", linestyle="--", linewidth=1.2)
    ax.text(len(labels) - 0.4, 12.6, "精品常見上限 12×", fontsize=7.5, color="#C9A227",
            ha="right")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels, fontsize=7)
    ax.set_ylabel("零售價 ÷ 出廠成本", fontsize=9)
    ax.set_title("定價倍率合理性檢查：原提案 vs 建議", fontsize=11, color="#1F3864",
                 fontweight="bold", pad=10)
    ax.legend(fontsize=8, frameon=False, ncol=2, loc="upper center",
              bbox_to_anchor=(0.5, -0.18))
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color="#E7E7E7", linewidth=0.8)
    ax.set_axisbelow(True)
    fig.savefig(CHARTS / "05_pricing.png")
    plt.close(fig)
    return CHARTS / "05_pricing.png"


def chart_pnl(pnl: dict, payback: dict) -> Path:
    keys = ["y1_original", "y1_suggested", "y2_original", "y2_suggested"]
    labels = ["第 1 年\n原提案價", "第 1 年\n建議價", "第 2 年\n原提案價", "第 2 年\n建議價"]
    rev = [pnl[k]["net_revenue"] for k in keys]
    prof = [pnl[k]["net_profit"] for k in keys]

    fig, ax = plt.subplots(figsize=(8.6, 3.9))
    x = range(len(keys))
    ax.bar([i - 0.2 for i in x], [v / 10000 for v in rev], width=0.38,
           color="#2E5496", label="品牌實收營收（萬）")
    ax.bar([i + 0.2 for i in x], [v / 10000 for v in prof], width=0.38,
           color="#548235", label="稅後淨利（萬）")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels, fontsize=8.5)
    ax.set_ylabel("TWD 萬元", fontsize=9)
    ax.set_title("損益情境比較（四種組合）", fontsize=11, color="#1F3864",
                 fontweight="bold", pad=10)
    ax.legend(fontsize=8, frameon=False, ncol=2, loc="upper center",
              bbox_to_anchor=(0.5, -0.16))
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color="#E7E7E7", linewidth=0.8)
    ax.set_axisbelow(True)
    for i, v in enumerate(prof):
        ax.text(i + 0.2, v / 10000 + 8, f"{v/10000:,.0f}", ha="center", fontsize=7.5,
                color="#548235", fontweight="bold")
    fig.savefig(CHARTS / "06_pnl.png")
    plt.close(fig)
    return CHARTS / "06_pnl.png"


# ── 主流程 ────────────────────────────────────────────────────────────────
def load():
    fin = json.loads((DATA / "finance_model.json").read_text(encoding="utf-8"))
    products = json.loads((DATA / "products.json").read_text(encoding="utf-8"))
    prices = json.loads((DATA / "ingredient_prices.json").read_text(encoding="utf-8"))
    with (DATA / "cost_report.csv").open(encoding="utf-8-sig") as f:
        cost_rows = list(csv.DictReader(f))
    with (DATA / "cost_detail.csv").open(encoding="utf-8-sig") as f:
        detail = list(csv.DictReader(f))
    return fin, products, prices, cost_rows, detail


def build():
    CHARTS.mkdir(parents=True, exist_ok=True)
    setup_mpl()

    fin, products, prices, cost_rows, detail = load()
    cost = fin["cost"]
    inv = fin["investment"]
    fac = fin["factory"]
    ai = fin["ai_team"]
    pricing = fin["pricing"]
    pnl = fin["pnl"]
    payback = fin["payback"]
    wf = fin["waterfall"][0]
    asm = fin["assumptions"]

    plist = products["products"] if isinstance(products, dict) and "products" in products else products
    pmap = {p["sku_code"]: p for p in plist}

    # 抓價統計
    ok = sum(1 for v in prices.get("prices", prices).values()
             if isinstance(v, dict) and v.get("status") == "ok") if isinstance(prices, dict) else 0
    total_ing = len(prices.get("prices", prices)) if isinstance(prices, dict) else 0
    l1 = sum(int(r["l1_count"]) for r in cost_rows)
    l2 = sum(int(r["l2_count"]) for r in cost_rows)
    l3 = sum(int(r["l3_count"]) for r in cost_rows)
    fb = sum(int(r["fallback_count"]) for r in cost_rows)

    # 圖表
    p_cost = chart_cost_structure(cost)
    p_inv = chart_investment(inv["by_category"])
    p_ai = chart_ai(ai)
    p_wf = chart_waterfall(wf)
    p_price = chart_pricing(pricing)
    p_pnl = chart_pnl(pnl, payback)

    doc = Document()

    # 預設樣式
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.paragraph_format.line_spacing = 1.45
    st.paragraph_format.space_after = Pt(6)

    for sname, ssize in (("Heading 1", 19), ("Heading 2", 13), ("Heading 3", 11)):
        try:
            hs = doc.styles[sname]
        except KeyError:
            continue
        hs.font.name = HEAD_FONT
        hs.font.size = Pt(ssize)
        hs.font.bold = True
        hs.font.color.rgb = rgb(NAVY)
        rPr = hs.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), HEAD_FONT)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rFonts.get(qn(attr)) is not None:
                del rFonts.attrib[qn(attr)]

    # 目錄樣式（Word 更新 TOC 時會套用這兩個樣式）
    from docx.enum.style import WD_STYLE_TYPE

    for sname, indent, ssize in (("TOC 1", 0.0, 10.5), ("TOC 2", 0.7, 9.5)):
        try:
            ts = doc.styles[sname]
        except KeyError:
            ts = doc.styles.add_style(sname, WD_STYLE_TYPE.PARAGRAPH)
        ts.base_style = doc.styles["Normal"]
        ts.font.name = BODY_FONT
        ts.font.size = Pt(ssize)
        ts.font.bold = sname == "TOC 1"
        ts.paragraph_format.space_before = Pt(1.5)
        ts.paragraph_format.space_after = Pt(1.5)
        ts.paragraph_format.line_spacing = 1.0
        ts.paragraph_format.left_indent = Cm(indent)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.9)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(0.9)

    # 頁首
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(hp.add_run("Rokaizumi 合同會社　｜　頂級化妝品產品線　｜　投資與獲利分析報告"),
            size=8, color=GREY_D)
    add_pbdr(hp, GOLD, 6, 2)

    # 頁尾
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(fp.add_run("內部文件．未經林總核定不得外流　　—　　"), size=8, color=GREY_D)
    add_field(fp, "PAGE")
    set_run(fp.add_run(" / "), size=8, color=GREY_D)
    add_field(fp, "NUMPAGES")

    # 讓 Word 開啟時自動更新目錄
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)

    # ══════════════════════════════════════════════════════════════════
    # 封面
    # ══════════════════════════════════════════════════════════════════
    doc.add_paragraph().paragraph_format.space_after = Pt(30)

    cover = doc.add_table(rows=1, cols=1)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(cover, color=NAVY, sz=8)
    cc = cover.rows[0].cells[0]
    cc.width = Cm(16.4)
    cc.text = ""
    shade(cc, NAVY)

    p = cc.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run("ROKAIZUMI ATELIER"), size=11, bold=True, color=GOLD, font=HEAD_FONT)

    p = cc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    set_run(p.add_run("頂 級 化 妝 品 產 品 線"), size=26, bold=True, color=WHITE, font=HEAD_FONT)

    p = cc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run("投資・成本・AI Agent 團隊・獲利結構"), size=14, bold=True,
            color=GOLD_L, font=HEAD_FONT)

    p = cc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(22)
    set_run(p.add_run("完 整 分 析 報 告"), size=12, color="C5D0E6", font=HEAD_FONT)

    p = cc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run("14 支 SKU　｜　248 筆原料明細　｜　6 個 AI 值班入口"),
            size=10, color=WHITE, font=HEAD_FONT)

    p = cc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(26)
    set_run(p.add_run("STO＋／RWA 架構對映"), size=10, color=GOLD_L, font=HEAD_FONT)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    info_rows = [
        ["文件名稱", "頂級化妝品產品線：投資與獲利分析報告"],
        ["版本", "v1.0"],
        ["產出日期", datetime.now().strftime("%Y-%m-%d")],
        ["主體", "日本 Rokaizumi 合同會社"],
        ["委託製造", fac["name"]],
        ["資料基準", "finance_model.json／cost_report.csv／cost_detail.csv"],
        ["文件性質", "內部決策參考．非對外報價文件"],
        ["核定狀態", "未核定（所有能力 enabled = false）"],
    ]
    add_table(doc, ["項目", "內容"], info_rows, [3.6, 12.8],
              aligns=["left", "left"], header_bg=NAVY2, size=9.5, zebra=True,
              first_col_bold=True)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    callout(doc, "閱讀前必讀：本報告的數字性質",
            "本報告所有成本為「估算值」，由公開網路零售價＋參考價推算，"
            "不是代工廠正式報價。銷量為情境假設，不是市場調查結果。"
            "正式投資決策前，必須以安成生化的正式 RFQ 回價與實際預購數據替換。",
            bg=RED_L, bar=RED, title_color=RED)

    # ══════════════════════════════════════════════════════════════════
    # 目錄
    # ══════════════════════════════════════════════════════════════════
    h1_plain(doc, "目　錄", new_page=True)
    toc_p = doc.add_paragraph()
    add_field(toc_p, r'TOC \o "1-1" \h \z \u', size=10, color="333333")
    para(doc, "（若目錄未顯示，請在 Word 中按 Ctrl+A 後按 F9，選擇「更新整個目錄」。）",
         size=8.5, color=GREY_D, italic=True)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    h2_plain(doc, "本報告要回答的四個問題")
    q_rows = [
        ["1", "這條產品線要花多少錢？", "一次性投資 NT$ 616.5 萬，其中 20.2% 是產品開發、20.2% 是首批量產"],
        ["2", "每支產品的成本是多少？", "出廠成本 NT$ 74～375／支，總成本 NT$ 80～405／支"],
        ["3", "AI Agent 團隊能省多少？", "年省 NT$ 81.1 萬（42.4%），人力從 4 人降到 2 人"],
        ["4", "錢怎麼分？多久回本？", "原提案價 0.57 年回本；建議價 2.6 年回本，但定價可持續"],
    ]
    add_table(doc, ["#", "問題", "本報告的答案"], q_rows, [1.0, 5.4, 10.0],
              aligns=["center", "left", "left"], size=9, zebra=True)

    # ══════════════════════════════════════════════════════════════════
    # 第 1 章
    # ══════════════════════════════════════════════════════════════════
    h1(doc, "專案總覽", "第 1 章", new_page=True)

    para(doc, "本章說明這條產品線的範圍、結構與現況。所有數字都來自同一份資料檔，"
              "與網頁版、CSV 版完全一致。", align="justify")

    h2(doc, "1.1 關鍵數字")
    kpi_rows = [
        ["產品支數", "14 支", "4 支 P0 首發、6 支 P1、4 支 P2"],
        ["原料明細", f"{len(detail)} 筆", "含比例、功能分類、法規限量"],
        ["一次性投資", f"NT$ {n(inv['total_twd'])}", f"約 {inv['total_twd']/10000:,.0f} 萬元，7 大類 21 項"],
        ["第 1 年品牌實收", f"NT$ {n(pnl['y1_original']['net_revenue'])}", "原提案價情境"],
        ["第 1 年稅後淨利", f"NT$ {n(pnl['y1_original']['net_profit'])}", f"淨利率 {pct(pnl['y1_original']['net_margin_pct'])}"],
        ["回收期（原提案價）", f"{payback['y1_original_years']:.2f} 年", "約 7 個月"],
        ["回收期（建議價）", f"{payback['y1_suggested_years']:.2f} 年", "定價下修後的代價"],
        ["AI 團隊年省", f"NT$ {n(ai['summary']['saving_year_twd'])}", f"節省 {pct(ai['summary']['saving_pct'])}"],
    ]
    add_table(doc, ["指標", "數值", "說明"], kpi_rows, [4.2, 4.4, 7.8],
              aligns=["left", "right", "left"], size=9.5, zebra=True,
              first_col_bold=True)

    h2(doc, "1.2 產品線結構")
    para(doc, "14 支 SKU 依功能分為五個系列，代碼規則為 RA-系列字母＋序號。"
              "P0 為首批打樣與量產對象，P1 為第二波，P2 為形象與禮盒款。", align="justify")

    sku_rows, sku_bg, sku_fg = [], [], []
    for r in cost_rows:
        code = r["sku_code"]
        pr = r["priority"]
        bg, fg = PRIORITY_STYLE.get(pr, (None, None))
        sku_rows.append([
            code, r["name_zh"], r["form"], f"{r['size_ml']}",
            pr, n(r["exfactory_twd"], 1), n(r["total_cost_twd"], 1),
            n(r["suggested_retail_twd"]),
        ])
        sku_bg.append(bg)
        sku_fg.append(fg)
    add_table(doc, ["代碼", "品名", "劑型", "容量", "優先", "出廠成本", "總成本", "建議售價"],
              sku_rows, [1.8, 3.8, 3.2, 1.4, 1.0, 1.7, 1.5, 1.8],
              aligns=["center", "left", "left", "center", "center", "right", "right", "right"],
              size=8.5, row_bg=sku_bg, row_fg=sku_fg, header_size=8.5)
    caption(doc, "表 1-1　14 支 SKU 主檔（容量單位 mL；底色：紅＝P0 首發、橘＝P1、灰＝P2；成本單位 TWD）")

    callout(doc, "為什麼不做「一模一樣」的照抄",
            "原始需求是照抄 tenlead.com（天麗生技／Joli 系列）的品名與成分比例。"
            "這會同時觸犯商標法第 68／95 條、著作權法第 91 條、營業秘密法第 10／13-1 條、"
            "公平交易法第 21／22／25 條，代工廠還會依民法第 185 條負連帶責任。"
            "本專案改走「品類對標＋原創命名＋公開知識配方架構＋正式 RFQ」的合法路徑，"
            "市場定位與功效訴求對齊，但文字、名稱、配方比例全部原創。",
            bg=GOLD_L, bar=GOLD, title_color="7F6000")

    # ══════════════════════════════════════════════════════════════════
    # 第 2 章
    # ══════════════════════════════════════════════════════════════════
    h1(doc, "成本分析", "第 2 章", new_page=True)

    h2(doc, "2.1 成本模型")
    para(doc, "成本由四塊組成，逐層加上代工廠與品牌端的管銷：", align="justify")
    bullet(doc, "料體：Σ（配方比例 × 原料單價）× 匯率 × 充填重量 ÷（1 − 損耗率）")
    bullet(doc, "包材：瓶器、外盒、印刷、貼標，逐支指定")
    bullet(doc, "充填：產線人工與設備攤提，逐支指定")
    bullet(doc, "檢驗：微生物、重金屬、安定性抽驗攤提")
    para(doc, "出廠成本 =（料體＋包材＋充填＋檢驗）×（1＋代工廠管銷毛利 20%）；"
              "總成本 = 出廠成本 ×（1＋品牌端物流管銷 8%）。", size=9.5, color=GREY_D)

    h2(doc, "2.2 逐支成本明細")
    cost_rows_tbl = []
    for r in cost_rows:
        cost_rows_tbl.append([
            r["sku_code"], r["name_zh"],
            n(r["raw_material_twd"], 1), n(r["packaging_twd"], 1),
            n(r["filling_twd"], 1), n(r["qc_twd"], 1),
            n(r["exfactory_twd"], 1), n(r["total_cost_twd"], 1),
            pct(r["raw_cost_ratio_pct"]),
        ])
    add_table(doc, ["代碼", "品名", "料體", "包材", "充填", "檢驗", "出廠成本", "總成本", "料體%"],
              cost_rows_tbl, [1.8, 4.4, 1.4, 1.4, 1.3, 1.2, 1.7, 1.5, 1.7],
              aligns=["center", "left", "right", "right", "right", "right", "right", "right", "right"],
              size=8.5, zebra=True, header_size=8.5)
    caption(doc, "表 2-1　逐支成本結構（單位 TWD／支；料體占比＝料體 ÷ 總成本）")

    add_picture(doc, p_cost, 16.2)
    caption(doc, "圖 2-1　14 支 SKU 成本結構堆疊圖")

    h2(doc, "2.3 三個關鍵發現")
    callout(doc, "發現一：包材比料體更貴",
            "14 支中有 12 支的包材成本高於料體成本（僅 RA-L07 例外，RA-C01 打平）。"
            "這代表議價重點不在原料，而在瓶器、外盒與印刷。"
            "把公版瓶器換成私模、把燙金改成網印，省下的錢比壓原料價多得多。",
            bg=BLU_L, bar=NAVY2)
    callout(doc, "發現二：零售倍率普遍過高",
            "多數 SKU 的零售價 ÷ 出廠成本超過 20 倍，最高 RA-N09 達 36.2 倍。"
            "國際精品保養品的常見倍率是 6～12 倍。倍率過高在沒有品牌力支撐時，"
            "會直接反映在退貨率與比價流失上。",
            bg=ORG_L, bar=ORG, title_color="843C0C")
    callout(doc, "發現三：成本槓桿在包材與批量",
            "原料價格上漲 30%，對總成本的影響只有 2～15%。"
            "真正的成本槓桿是包材規格與生產批量。"
            "把 MOQ 從 3,000 拉到 10,000，包材單價通常可降 25～35%。",
            bg=GRN_L, bar=GRN, title_color="375623")

    h2(doc, "2.4 價格來源與可信度")
    para(doc, "成本引擎採三層價格回退，確保每一筆原料都有價格可用：", align="justify")
    layer_rows = [
        ["L1 自動抓取", f"{l1} 筆", "makingcosmetics.com 公開零售價", "高（有實際頁面）", "零售／小批量價，通常偏高 2～5 倍"],
        ["L2 參考價", f"{l2} 筆", "ingredient_master.json 內建參考單價", "中（產業經驗值）", "未逐項驗證"],
        ["L3 人工覆寫", f"{l3} 筆", "price_overrides.json", "最高（人工確認）", "目前為空，待代工廠報價填入"],
        ["Fallback", f"{fb} 筆", "類別中位數", "低", "僅為避免計算中斷"],
    ]
    add_table(doc, ["層級", "筆數", "來源", "可信度", "限制"], layer_rows,
              [2.6, 1.6, 4.6, 3.0, 4.6],
              aligns=["left", "right", "left", "left", "left"], size=8.5, zebra=True,
              first_col_bold=True)
    caption(doc, "表 2-2　三層價格回退機制")

    para(doc, f"自動抓價覆蓋率：{ok}／{total_ing} 項（{ok/total_ing*100:.0f}%）。"
              "胜肽、防曬劑、神經醯胺、植物油脂、基礎化學品命中率偏低，"
              "原因是這些品項在零售網站上不以原料形式販售，或名稱差異過大。"
              "台灣本地化工行（第一化工、城乙化工）多數沒有公開價格或網站不可達。",
         size=9.5, color=GREY_D, align="justify")

    # ══════════════════════════════════════════════════════════════════
    # 第 3 章
    # ══════════════════════════════════════════════════════════════════
    h1(doc, "定價合理性", "第 3 章", new_page=True)

    para(doc, "本章把「原提案價」與「建議價」並列，讓倍率問題一眼看得出來。"
              "建議價的算法是：出廠成本 × 7 倍，再向上取整到 10 元。", align="justify")

    price_rows, price_bg, price_fg = [], [], []
    for o, s in zip(pricing["original"], pricing["suggested"]):
        m = o["multiple"]
        if m >= 25:
            bg, fg = RED_L, RED
        elif m >= 15:
            bg, fg = ORG_L, ORG
        else:
            bg, fg = GRN_L, GRN
        price_rows.append([
            o["sku_code"], o["name_zh"], n(o["exfactory"], 1),
            n(o["price"]), f"{m:.1f}×", n(s["price"]), f"{s['multiple']:.1f}×",
            n(s["price"] - o["price"]),
        ])
        price_bg.append(bg)
        price_fg.append(fg)
    add_table(doc, ["代碼", "品名", "出廠成本", "原提案價", "倍率", "建議價", "倍率", "價差"],
              price_rows, [1.8, 5.4, 1.8, 1.7, 1.3, 1.6, 1.2, 1.6],
              aligns=["center", "left", "right", "right", "center", "right", "center", "right"],
              size=8.5, row_bg=price_bg, row_fg=price_fg, header_size=8.5)
    caption(doc, "表 3-1　定價倍率對照（底色：紅＝25 倍以上、橘＝15～25 倍、綠＝15 倍以下）")

    add_picture(doc, p_price, 16.2)
    caption(doc, "圖 3-1　定價倍率合理性檢查")

    h2(doc, "3.1 該不該下修？")
    para(doc, "這是一個真正的兩難，本報告不替林總決定，但把代價算清楚：", align="justify")
    trade_rows = [
        ["維持原提案價", f"NT$ {n(pnl['y1_original']['net_profit'])}",
         f"{payback['y1_original_years']:.2f} 年",
         "回本極快，但倍率 15～36 倍在無品牌力時難以支撐，退貨與比價流失風險高"],
        ["下修至 7 倍", f"NT$ {n(pnl['y1_suggested']['net_profit'])}",
         f"{payback['y1_suggested_years']:.2f} 年",
         "回本變慢 4.6 倍，但價格站得住，適合長期品牌經營與代理制度"],
    ]
    add_table(doc, ["情境", "第 1 年稅後淨利", "回收期", "評語"], trade_rows,
              [3.0, 3.4, 2.0, 8.0],
              aligns=["left", "right", "center", "left"], size=9, zebra=True,
              first_col_bold=True)
    caption(doc, "表 3-2　兩種定價策略的代價")

    callout(doc, "本報告的建議（供林總核定）",
            "採「雙軌定價」：核心 4 支 P0 走 7～9 倍的可持續價，建立回購與代理基礎；"
            "形象款（RA-R12 禮盒、RA-N10 A 醇晚霜）維持高倍率，作為價格錨點。"
            "這樣既保住品牌高度，也不會讓主力商品因倍率過高而流失。",
            bg=GOLD_L, bar=GOLD, title_color="7F6000")

    # ══════════════════════════════════════════════════════════════════
    # 第 4 章
    # ══════════════════════════════════════════════════════════════════
    h1(doc, "化工廠投資與合作", "第 4 章", new_page=True)

    h2(doc, "4.1 合作定位")
    para(doc, f"代工廠：{fac['name']}。角色：{fac['role']}。", bold=True)
    callout(doc, "核心原則：代工廠出資 0 元、不佔股",
            fac["factory_note"] + " 這樣做的好處是配方歸屬乾淨、股權結構單純、"
            "未來換廠或增加第二供應商時沒有牽絆。代價是品牌方要自己扛全部開發費用與庫存風險。",
            bg=BLU_L, bar=NAVY2)

    h2(doc, "4.2 代工廠提供什麼")
    fac_rows = [[a, b, c] for a, b, c in fac["factory_contribution"]]
    add_table(doc, ["投入項目", "內容", "條件"], fac_rows, [3.0, 7.4, 6.0],
              aligns=["left", "left", "left"], size=9, zebra=True, first_col_bold=True)
    caption(doc, "表 4-1　代工廠投入")

    h2(doc, "4.3 品牌方負擔的費用")
    pay_rows = [[a, b, c] for a, b, c in fac["brand_side_pays"]]
    add_table(doc, ["費用項目", "單價", "說明"], pay_rows, [3.0, 4.4, 9.0],
              aligns=["left", "left", "left"], size=9, zebra=True, first_col_bold=True)
    caption(doc, "表 4-2　品牌方負擔")

    h2(doc, "4.4 合約必要條件")
    para(doc, "以下八條是簽約前必須白紙黑字寫進合約的條件，缺一條就不簽：", align="justify")
    ct_rows = [[str(i + 1), a, b, c] for i, (a, b, c) in enumerate(fac["contract_terms"])]
    add_table(doc, ["#", "條款", "約定內容", "為什麼重要"], ct_rows, [0.9, 3.0, 6.4, 6.1],
              aligns=["center", "left", "left", "left"], size=8.5, zebra=True)
    caption(doc, "表 4-3　合約必要條件")

    h2(doc, "4.5 風險與對策")
    risk_rows = []
    for i, r in enumerate(fac["risk_notes"]):
        risk_rows.append([str(i + 1), r])
    add_table(doc, ["#", "風險提示"], risk_rows, [0.9, 15.5],
              aligns=["center", "left"], size=9, zebra=True)
    caption(doc, "表 4-4　風險提示")

    h2(doc, "4.6 投資明細")
    add_picture(doc, p_inv, 15.0)
    caption(doc, "圖 4-1　一次性投資結構")

    inv_rows = []
    for r in inv["rows"]:
        inv_rows.append([r["cat_code"], r["category"], r["item"],
                         n(r["amount_twd"]), r["note"]])
    inv_rows.append(["", "合計", "", n(inv["total_twd"]), ""])
    add_table(doc, ["類", "類別", "項目", "金額 TWD", "說明"], inv_rows,
              [0.9, 2.6, 4.4, 2.4, 6.1],
              aligns=["center", "left", "left", "right", "left"], size=8,
              zebra=True, bold_rows=(len(inv_rows) - 1,))
    caption(doc, "表 4-5　一次性投資明細（21 項）")

    # ══════════════════════════════════════════════════════════════════
    # 第 5 章
    # ══════════════════════════════════════════════════════════════════
    h1(doc, "AI Agent 團隊", "第 5 章", new_page=True)

    para(doc, "20 個 Agent 合併成 6 個值班入口（E1～E6），由 2 位真人負責判斷、核定與對外關係。"
              "AI 不取代人，是讓人從 4 人降到 2 人。", align="justify")

    h2(doc, "5.1 六個值班入口")
    ai_rows = []
    for r in ai["rows"]:
        ai_rows.append([r["code"], r["name"], r["agents"], r["human_role"],
                        n(r["human_cost_twd"]), n(r["ai_cost_twd"], 0), r["duty"]])
    ai_rows.append(["", "合計", "", f"{ai['summary']['full_human_headcount']} 人",
                    n(ai["summary"]["full_human_month_twd"]),
                    n(sum(r["ai_cost_twd"] for r in ai["rows"]), 0), ""])
    add_table(doc, ["入口", "職能", "對應 Agent", "替代人力", "人力月成本", "AI 月成本", "主要職責"],
              ai_rows, [1.2, 2.2, 2.6, 2.4, 1.9, 1.7, 4.4],
              aligns=["center", "left", "left", "left", "right", "right", "left"],
              size=8, zebra=True, bold_rows=(len(ai_rows) - 1,), header_size=8)
    caption(doc, "表 5-1　六個值班入口（成本單位 TWD／月）")

    add_picture(doc, p_ai, 16.0)
    caption(doc, "圖 5-1　人力成本 vs AI Agent 成本")

    h2(doc, "5.2 成本對比")
    s = ai["summary"]
    cmp_rows = [
        ["純人力（4 人）", f"{s['full_human_headcount']} 人",
         f"NT$ {n(s['full_human_month_twd'])}", f"NT$ {n(s['full_human_year_twd'])}", "—"],
        ["精簡人力（2 人）", f"{s['lean_headcount']} 人",
         f"NT$ {n(s['lean_human_year_twd']/12)}", f"NT$ {n(s['lean_human_year_twd'])}", "—"],
        ["AI Agent 維運", "—", f"NT$ {n(s['ai_month_twd'])}",
         f"NT$ {n(s['ai_year_twd'])}", f"USD {s['ai_month_usd']}／月"],
        ["精簡後合計", f"{s['lean_headcount']} 人 + AI", "—",
         f"NT$ {n(s['lean_total_year_twd'])}", "—"],
        ["年節省", "—", "—", f"NT$ {n(s['saving_year_twd'])}", f"節省 {pct(s['saving_pct'])}"],
    ]
    add_table(doc, ["方案", "人力", "月成本", "年成本", "備註"], cmp_rows,
              [3.4, 2.6, 3.0, 3.4, 4.0],
              aligns=["left", "center", "right", "right", "left"], size=9,
              zebra=True, bold_rows=(3, 4))
    caption(doc, "表 5-2　人力與 AI 成本對比（TWD）")

    callout(doc, "AI 團隊的邊界",
            s["note"] + " AI 只做「整理、比對、草擬、提醒」，"
            "不做「公開發布、付費、簽約、送件、正式發行、資料外傳、刪除正式資料」。"
            "這七類高影響操作一律進人工核定佇列，由林總決定。",
            bg=GRN_L, bar=GRN, title_color="375623")

    # ══════════════════════════════════════════════════════════════════
    # 第 6 章
    # ══════════════════════════════════════════════════════════════════
    h1(doc, "獲利比例與回收", "第 6 章", new_page=True)

    h2(doc, "6.1 錢怎麼分：價值鏈拆解")
    para(doc, f"以 {wf['sku_code']} {wf['name_zh']} 為例，建議零售價 NT$ {n(wf['retail_price'])}，"
              "從零售價一路拆到品牌稅後淨利：", align="justify")

    wf_rows = []
    for it in wf["items"]:
        wf_rows.append([it["name"], n(it["amount"], 1), pct(it["pct"]), it["note"]])
    add_table(doc, ["項目", "TWD／支", "占零售價", "說明"], wf_rows,
              [3.4, 2.2, 2.0, 8.8],
              aligns=["left", "right", "right", "left"], size=8.5, zebra=True,
              bold_rows=(len(wf_rows) - 1,))
    caption(doc, "表 6-1　價值鏈拆解（單支）")

    add_picture(doc, p_wf, 16.0)
    caption(doc, "圖 6-1　價值鏈瀑布圖")

    h2(doc, "6.2 損益表")
    para(doc, "四種情境並列：第 1 年／第 2 年 × 原提案價／建議價。", align="justify")

    keys = [("y1_original", "第 1 年 原提案價"), ("y1_suggested", "第 1 年 建議價"),
            ("y2_original", "第 2 年 原提案價"), ("y2_suggested", "第 2 年 建議價")]
    exp_keys = list(pnl["y1_original"]["expenses"].keys())

    pl_rows = [["零售總額（含稅）"] + [n(pnl[k]["gross_retail_revenue"]) for k, _ in keys]]
    pl_rows.append(["品牌實收營收"] + [n(pnl[k]["net_revenue"]) for k, _ in keys])
    for ek in exp_keys:
        pl_rows.append([ek] + [n(pnl[k]["expenses"].get(ek, 0)) for k, _ in keys])
    pl_rows.append(["費用合計"] + [n(pnl[k]["total_expenses"]) for k, _ in keys])
    pl_rows.append(["稅前淨利"] + [n(pnl[k]["ebt"]) for k, _ in keys])
    pl_rows.append(["營所稅"] + [n(pnl[k]["tax"]) for k, _ in keys])
    pl_rows.append(["稅後淨利"] + [n(pnl[k]["net_profit"]) for k, _ in keys])
    pl_rows.append(["淨利率"] + [pct(pnl[k]["net_margin_pct"]) for k, _ in keys])

    add_table(doc, ["項目"] + [lbl for _, lbl in keys], pl_rows,
              [4.4, 3.0, 3.0, 3.0, 3.0],
              aligns=["left", "right", "right", "right", "right"], size=8.5,
              zebra=True, bold_rows=(0, 1, len(pl_rows) - 4, len(pl_rows) - 1))
    caption(doc, "表 6-2　損益表（TWD）")

    add_picture(doc, p_pnl, 16.0)
    caption(doc, "圖 6-2　損益情境比較")

    h2(doc, "6.3 回收期")
    pb_rows = [
        ["第 1 年 原提案價", f"{payback['y1_original_years']:.2f} 年",
         f"約 {payback['y1_original_years']*12:.1f} 個月"],
        ["第 1 年 建議價", f"{payback['y1_suggested_years']:.2f} 年",
         f"約 {payback['y1_suggested_years']*12:.1f} 個月"],
        ["第 2 年 原提案價", f"{payback['y2_original_years']:.2f} 年",
         f"約 {payback['y2_original_years']*12:.1f} 個月"],
        ["第 2 年 建議價", f"{payback['y2_suggested_years']:.2f} 年",
         f"約 {payback['y2_suggested_years']*12:.1f} 個月"],
    ]
    add_table(doc, ["情境", "回收期（年）", "換算"], pb_rows, [5.0, 4.0, 4.0],
              aligns=["left", "right", "right"], size=9.5, zebra=True,
              first_col_bold=True)
    caption(doc, f"表 6-3　回收期（一次性投資 NT$ {n(inv['total_twd'])}）")

    h2(doc, "6.4 財務假設")
    asm_rows = [
        ["匯率", f"{asm['usd_twd']:.1f} TWD／USD", "近期均值"],
        ["營所稅率", pct(asm["tax_rate"] * 100, 0), "台灣營利事業所得稅"],
        ["通路抽成", pct(asm["channel_cut"] * 100, 0), "百貨／藥粧"],
        ["自營比重", pct(asm["self_direct_share"] * 100, 0), "其餘走通路"],
        ["金流手續費", pct(asm["payment_fee"] * 100, 0), "佔自營營收"],
        ["物流倉儲", pct(asm["logistics_pct"] * 100, 0), "佔品牌實收"],
        ["行銷", pct(asm["marketing_pct"] * 100, 0), "佔品牌實收"],
        ["代理佣金", pct(asm["agent_commission_pct"] * 100, 0), "佔通路營收"],
        ["社群貢獻回饋", pct(asm["community_pct"] * 100, 0), "佔品牌實收，與佣金分開記帳"],
        ["雜支", pct(asm["misc_pct"] * 100, 0), "佔品牌實收"],
        ["人事", f"{asm['headcount']} 人 × NT$ {n(asm['salary_per_month'])}／月", "精簡後編制"],
        ["AI 預算", f"USD {asm['ai_budget_usd_month']}／月", "6 個入口共用"],
    ]
    add_table(doc, ["假設", "數值", "說明"], asm_rows, [3.4, 5.4, 7.6],
              aligns=["left", "left", "left"], size=9, zebra=True, first_col_bold=True)
    caption(doc, "表 6-4　財務模型假設")

    callout(doc, "必須誠實說明的一件事",
            "本報告的銷量假設（第 1 年 4 支 P0 共 16,800 支、第 2 年 14 支全開）"
            "是憑空設定的情境，不是市場調查結果，也不是預購數據。"
            "回收期 0.57 年這個數字看起來很漂亮，但它完全建立在銷量假設上。"
            "銷量只要掉一半，回收期就變 1.1 年；掉到三成，就變 1.9 年。",
            bg=RED_L, bar=RED, title_color=RED)

    # ══════════════════════════════════════════════════════════════════
    # 第 7 章
    # ══════════════════════════════════════════════════════════════════
    h1(doc, "法遵邊界與 STO＋／RWA", "第 7 章", new_page=True)

    h2(doc, "7.1 法遵紅線")
    law_rows = [
        ["商標法 68、95 條", "不得使用相同或近似他人註冊商標於同一或類似商品", "品名、系列名全部原創，並先做商標檢索"],
        ["著作權法 91 條", "不得重製他人語文、美術著作", "文案、成分表敘述、視覺全部自行撰寫"],
        ["營業秘密法 10、13-1 條", "不得取得、使用他人營業秘密", "不取得對方配方，改以公開知識重建"],
        ["公平交易法 21、22、25 條", "不得為虛偽不實或引人錯誤之表示", "功效宣稱須有依據，不得攀附他人商譽"],
        ["化粧品衛生安全管理法 7、8、22 條", "成分限量、標示、登錄、PIF 文件", "逐項核對法規限量，建立 PIF"],
        ["民法 185 條", "共同侵權連帶賠償", "代工廠若照抄也會被連帶求償，故不接受照抄訂單"],
    ]
    add_table(doc, ["法規", "禁止事項", "本專案的對策"], law_rows, [3.6, 6.4, 6.4],
              aligns=["left", "left", "left"], size=8.5, zebra=True, first_col_bold=True)
    caption(doc, "表 7-1　法遵紅線與對策")

    h2(doc, "7.2 STO＋／RWA 對映")
    para(doc, "這條產品線不是獨立專案，而是掛在既有的 STO＋ 同一代幣多狀態模型下。"
              "實體商品對映為 RWA，消費者端的權益先設計為「不可轉讓、不承諾轉 ST」。", align="justify")
    sto_rows = [
        ["RWA 錨定", "每支 SKU 的批次生產紀錄、檢驗報告、庫存", "上鏈只放 merkle_root 與 leaf_count（D-014）"],
        ["多狀態模型", "原料 → 半成品 → 成品 → 已售出 → 已回購", "狀態轉換需對應實體事件"],
        ["消費者權益", "購買憑證、回購權、社群貢獻點數", "不可轉讓、不承諾轉 ST（D-017）"],
        ["代理制度", "單層推薦佣金", "先只做單層（D-021），與社群貢獻分開記帳"],
        ["市場範圍", "台灣＋日本先行", "第一階段面對消費者的市場（D-018）"],
        ["資料邊界", "照片、個資、健康資料", "一律不送境外 API（D-020）"],
        ["能力開關", "所有未核定能力", "預設 enabled = false"],
    ]
    add_table(doc, ["面向", "對映內容", "約束"], sto_rows, [2.8, 6.6, 7.0],
              aligns=["left", "left", "left"], size=8.5, zebra=True, first_col_bold=True)
    caption(doc, "表 7-2　STO＋／RWA 對映")

    # ══════════════════════════════════════════════════════════════════
    # 第 8 章
    # ══════════════════════════════════════════════════════════════════
    h1(doc, "待林總決策事項", "第 8 章", new_page=True)

    para(doc, "以下 10 項需要林總核定。在核定前，相關能力一律維持 enabled = false，"
              "不會自行執行任何高影響操作。", align="justify")

    dec_rows = [
        ["P-101", "品牌名是否採用 ROKAIZUMI ATELIER／露海澄研", "品牌", "高", "影響所有包裝與商標申請"],
        ["P-102", "是否先做商標檢索（台灣＋日本）", "法遵", "高", "未檢索就送件有駁回與侵權風險"],
        ["P-103", "首批打樣支數（建議 4 支 P0）", "產品", "中", "影響打樣費 24 萬"],
        ["P-104", "是否授權把 RFQ 寄給安成生化", "採購", "高", "屬對外送件，須林總核定"],
        ["P-105", "建議零售價是否照提案", "定價", "高", "直接決定損益與回收期"],
        ["P-106", "頭皮線是否與外泌體專案合流", "策略", "中", "影響 S13／S14 的定位"],
        ["P-107", "是否接受「估算值」作為現階段成本依據", "財務", "中", "不接受則須先取得代工廠報價"],
        ["P-108", "建議零售價是否下修至 6～8 倍區間", "定價", "高", "與 P-105 連動"],
        ["P-109", "是否授權寄 RFQ 以取得實價", "採購", "高", "與 P-104 連動"],
        ["P-110", "包材是否採公版先行", "包材", "中", "可省私模費 28 萬，但外觀較無差異化"],
    ]
    add_table(doc, ["編號", "決策事項", "類別", "優先", "影響"], dec_rows,
              [1.4, 6.4, 1.4, 1.2, 6.0],
              aligns=["center", "left", "center", "center", "left"], size=8.5, zebra=True)
    caption(doc, "表 8-1　待決策事項")

    h2(doc, "8.1 建議的下一步順序")
    step_rows = [
        ["1", "核定 P-101、P-102", "品牌名與商標檢索", "1 週"],
        ["2", "核定 P-104、P-109 並寄出 RFQ", "取得安成生化正式報價", "2～3 週"],
        ["3", "用正式報價替換估算值", "重跑 cost_engine 與 finance_model", "1 天"],
        ["4", "核定 P-105、P-108 定價", "確定損益與回收期", "1 週"],
        ["5", "核定 P-103 打樣", "4 支 P0 進入打樣", "6～8 週"],
        ["6", "核定 P-110 包材", "公版先行或開私模", "與打樣並行"],
    ]
    add_table(doc, ["#", "動作", "產出", "預估時間"], step_rows, [0.9, 5.6, 6.4, 3.5],
              aligns=["center", "left", "left", "center"], size=9, zebra=True)
    caption(doc, "表 8-2　建議執行順序")

    doc.add_paragraph().paragraph_format.space_after = Pt(14)
    callout(doc, "本報告的定位",
            "這是一份「決策前的估算報告」，不是「投資建議書」。"
            "它的價值在於把不確定性攤開來看：哪些數字有依據、哪些是假設、"
            "哪些必須等代工廠報價才能定案。"
            "所有數字都可以用 scripts\\run_cost_analysis.ps1 與 scripts\\finance_model.py 重跑驗證。",
            bg=BLU_L, bar=NAVY2)

    # 結尾（移除 callout 尾端空白段，避免多出一頁）
    last = doc.paragraphs[-1]
    if not last.text.strip():
        last._element.getparent().remove(last._element)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    set_run(p.add_run("— 報告結束 —"), size=10, color=GREY_D, font=HEAD_FONT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(f"Rokaizumi 合同會社　｜　{datetime.now().strftime('%Y-%m-%d')}　｜　v1.0"),
            size=8.5, color=GREY_D)

    doc.save(str(OUT_DOCX))
    return OUT_DOCX


if __name__ == "__main__":
    try:
        out = build()
    except Exception as exc:  # noqa: BLE001
        print(f"失敗：{exc}", file=sys.stderr)
        raise
    size = out.stat().st_size
    print(f"完成 → {out}  ({size:,} bytes)")
    print(f"圖表 → {CHARTS}")
